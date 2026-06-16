from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "presentation_assets"
SHOTS = ASSETS / "screenshots"
OUTPUT = ASSETS / "EcoSafe_Soutenance.pptx"
ICON = ROOT / "web" / "icons" / "Icon-512.png"


GREEN = RGBColor(28, 179, 116)
GREEN_DARK = RGBColor(18, 117, 76)
GREEN_SOFT = RGBColor(232, 247, 240)
BLUE = RGBColor(42, 115, 255)
BLUE_SOFT = RGBColor(233, 241, 255)
ORANGE = RGBColor(244, 144, 34)
ORANGE_SOFT = RGBColor(255, 243, 227)
DARK = RGBColor(27, 38, 59)
TEXT = RGBColor(62, 72, 89)
MUTED = RGBColor(107, 119, 140)
LIGHT = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
BORDER = RGBColor(221, 228, 237)
RED = RGBColor(230, 90, 90)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_box(slide, x, y, w, h, color, radius=True, line=None):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line or color
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=20,
    color=TEXT,
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name="Aptos",
):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, Inches(0.6), Inches(0.35), Inches(7.0), Inches(0.6), 28, DARK, True, font_name="Aptos Display")
    if subtitle:
        add_text(slide, subtitle, Inches(0.6), Inches(0.82), Inches(8.5), Inches(0.45), 12, MUTED)


def add_picture_contain(slide, path, x, y, w, h):
    img = Image.open(path)
    img_w, img_h = img.size
    target_ratio = w / h
    img_ratio = img_w / img_h
    if img_ratio > target_ratio:
        pic_w = w
        pic_h = w / img_ratio
        pic_x = x
        pic_y = y + (h - pic_h) / 2
    else:
        pic_h = h
        pic_w = h * img_ratio
        pic_x = x + (w - pic_w) / 2
        pic_y = y
    slide.shapes.add_picture(str(path), pic_x, pic_y, width=pic_w, height=pic_h)


def add_metric_card(slide, x, y, w, h, title, value, accent, note):
    add_box(slide, x, y, w, h, WHITE, line=BORDER)
    add_text(slide, title, x + Inches(0.18), y + Inches(0.12), w - Inches(0.3), Inches(0.2), 11, MUTED, True)
    add_text(slide, value, x + Inches(0.18), y + Inches(0.34), w - Inches(0.3), Inches(0.35), 22, accent, True, font_name="Aptos Display")
    add_text(slide, note, x + Inches(0.18), y + Inches(0.72), w - Inches(0.3), Inches(0.35), 10, TEXT)


def add_pill(slide, text, x, y, w, color, txt=WHITE):
    add_box(slide, x, y, w, Inches(0.34), color)
    add_text(slide, text, x, y + Inches(0.02), w, Inches(0.2), 10, txt, True, PP_ALIGN.CENTER)


def add_step(slide, number, title, desc, x, y, color):
    circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, Inches(0.42), Inches(0.42))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.color.rgb = color
    add_text(slide, str(number), x, y + Inches(0.05), Inches(0.42), Inches(0.22), 16, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, title, x + Inches(0.55), y - Inches(0.02), Inches(2.2), Inches(0.25), 13, DARK, True)
    add_text(slide, desc, x + Inches(0.55), y + Inches(0.22), Inches(2.5), Inches(0.45), 10, MUTED)


def build_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_box(slide, Inches(0), Inches(0), Inches(5.4), Inches(7.5), GREEN, radius=False)
    add_text(slide, "EcoSafe", Inches(0.7), Inches(1.0), Inches(3.3), Inches(0.6), 28, WHITE, True, font_name="Aptos Display")
    add_text(slide, "Soutenance du projet mobile", Inches(0.7), Inches(1.55), Inches(3.8), Inches(0.4), 18, WHITE, True)
    add_text(
        slide,
        "Une application qui compare les trajets selon la securite, le CO2 et les points gagnes.",
        Inches(0.7),
        Inches(2.1),
        Inches(3.8),
        Inches(0.9),
        14,
        WHITE,
    )
    add_pill(slide, "Equipe de 4", Inches(0.7), Inches(3.1), Inches(1.15), WHITE, GREEN_DARK)
    add_pill(slide, "Flutter + Firebase", Inches(1.95), Inches(3.1), Inches(1.65), WHITE, GREEN_DARK)
    add_pill(slide, "Securite + Ecologie", Inches(0.7), Inches(3.58), Inches(2.0), WHITE, GREEN_DARK)
    if ICON.exists():
        slide.shapes.add_picture(str(ICON), Inches(0.7), Inches(4.5), width=Inches(1.2), height=Inches(1.2))
    add_box(slide, Inches(5.0), Inches(0.45), Inches(7.7), Inches(6.6), LIGHT, line=LIGHT)
    add_picture_contain(slide, SHOTS / "01_carte.png", Inches(5.25), Inches(0.72), Inches(3.0), Inches(5.9))
    add_picture_contain(slide, SHOTS / "02_transport.png", Inches(8.45), Inches(1.1), Inches(3.0), Inches(5.55))


def build_problem_solution(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_title(slide, "Le Projet En Une Idee", "Une navigation plus utile pour les deplacements du quotidien.")
    add_box(slide, Inches(0.6), Inches(1.35), Inches(3.8), Inches(5.35), WHITE, line=BORDER)
    add_text(slide, "Le besoin", Inches(0.85), Inches(1.6), Inches(2.0), Inches(0.3), 20, DARK, True)
    add_metric_card(slide, Inches(0.85), Inches(2.1), Inches(3.3), Inches(1.05), "Probleme 1", "Quel trajet choisir ?", ORANGE, "Le plus rapide n'est pas toujours le plus sur.")
    add_metric_card(slide, Inches(0.85), Inches(3.3), Inches(3.3), Inches(1.05), "Probleme 2", "Impact CO2 peu visible", GREEN, "On veut rendre le cout environnemental concret.")
    add_metric_card(slide, Inches(0.85), Inches(4.5), Inches(3.3), Inches(1.05), "Probleme 3", "Peu de motivation", BLUE, "Les points et recompenses poussent aux bons choix.")

    add_box(slide, Inches(4.7), Inches(1.35), Inches(8.0), Inches(5.35), WHITE, line=BORDER)
    add_text(slide, "Notre reponse", Inches(5.0), Inches(1.6), Inches(2.5), Inches(0.3), 20, DARK, True)
    add_pill(slide, "Comparer", Inches(5.0), Inches(2.05), Inches(0.95), GREEN)
    add_pill(slide, "Mesurer", Inches(6.05), Inches(2.05), Inches(0.95), BLUE)
    add_pill(slide, "Recompenser", Inches(7.1), Inches(2.05), Inches(1.25), ORANGE)
    add_text(slide, "Carte interactive", Inches(5.0), Inches(2.55), Inches(2.1), Inches(0.28), 15, DARK, True)
    add_text(slide, "Plusieurs options de trajet a partir d'un depart et d'une arrivee.", Inches(5.0), Inches(2.82), Inches(2.9), Inches(0.5), 11, MUTED)
    add_text(slide, "Score securite", Inches(5.0), Inches(3.55), Inches(2.1), Inches(0.28), 15, DARK, True)
    add_text(slide, "L'eclairage, le trafic, la frequentation et les pistes cyclables sont combines.", Inches(5.0), Inches(3.82), Inches(3.1), Inches(0.55), 11, MUTED)
    add_text(slide, "Points & boutique", Inches(5.0), Inches(4.65), Inches(2.2), Inches(0.28), 15, DARK, True)
    add_text(slide, "Les choix eco et surs rapportent des points utilisables dans la boutique.", Inches(5.0), Inches(4.92), Inches(3.1), Inches(0.55), 11, MUTED)
    add_picture_contain(slide, SHOTS / "05_boutique.png", Inches(8.8), Inches(1.9), Inches(2.85), Inches(4.5))


def build_flow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_title(slide, "Fonctionnement", "Un parcours simple, de la recherche au gain de points.")
    add_step(slide, 1, "Choisir un trajet", "Depart, arrivee, mode de transport.", Inches(0.8), Inches(1.55), GREEN)
    add_step(slide, 2, "Comparer", "CO2, securite, duree et points.", Inches(3.45), Inches(1.55), BLUE)
    add_step(slide, 3, "Selectionner", "L'utilisateur choisit l'option la plus interessante.", Inches(6.1), Inches(1.55), ORANGE)
    add_step(slide, 4, "Cumuler", "Les points alimentent le profil et la boutique.", Inches(8.75), Inches(1.55), GREEN_DARK)
    add_picture_contain(slide, SHOTS / "01_carte.png", Inches(0.8), Inches(3.0), Inches(2.4), Inches(3.8))
    add_picture_contain(slide, SHOTS / "02_transport.png", Inches(3.6), Inches(3.0), Inches(2.4), Inches(3.8))
    add_picture_contain(slide, SHOTS / "05_boutique.png", Inches(8.8), Inches(3.0), Inches(2.4), Inches(3.8))
    add_box(slide, Inches(6.25), Inches(3.2), Inches(2.1), Inches(3.4), GREEN_SOFT, line=GREEN_SOFT)
    add_text(slide, "Resultat attendu", Inches(6.48), Inches(3.45), Inches(1.7), Inches(0.28), 16, GREEN_DARK, True)
    add_text(
        slide,
        "Un trajet plus clair a comprendre,\nplus responsable,\net plus engageant pour l'utilisateur.",
        Inches(6.48),
        Inches(4.0),
        Inches(1.55),
        Inches(1.6),
        13,
        TEXT,
        False,
        PP_ALIGN.CENTER,
    )


def build_pages(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_title(slide, "Pages Principales", "Des ecrans courts, visuels et centres sur la decision.")

    cards = [
        ("Carte", SHOTS / "01_carte.png", Inches(0.6)),
        ("Transport", SHOTS / "02_transport.png", Inches(3.2)),
        ("Boutique", SHOTS / "05_boutique.png", Inches(5.8)),
    ]
    for title, path, x in cards:
        add_box(slide, x, Inches(1.4), Inches(2.25), Inches(4.25), WHITE, line=BORDER)
        add_text(slide, title, x + Inches(0.15), Inches(1.55), Inches(1.0), Inches(0.25), 14, DARK, True)
        add_picture_contain(slide, path, x + Inches(0.12), Inches(1.85), Inches(2.0), Inches(3.55))

    add_box(slide, Inches(8.5), Inches(1.4), Inches(4.2), Inches(1.95), WHITE, line=BORDER)
    add_text(slide, "Module securite", Inches(8.75), Inches(1.68), Inches(2.0), Inches(0.25), 16, DARK, True)
    add_text(slide, "Score moyen\nCritere par critere\nAlertes meteo et trafic", Inches(8.75), Inches(2.1), Inches(1.8), Inches(0.9), 14, GREEN_DARK, False)

    add_box(slide, Inches(8.5), Inches(3.65), Inches(4.2), Inches(1.95), WHITE, line=BORDER)
    add_text(slide, "Profil & stats", Inches(8.75), Inches(3.93), Inches(2.0), Inches(0.25), 16, DARK, True)
    add_text(slide, "Points cumules\nCO2 economise\nSuivi hebdo et mensuel", Inches(8.75), Inches(4.35), Inches(1.8), Inches(0.9), 14, BLUE, False)

    add_box(slide, Inches(8.5), Inches(5.9), Inches(4.2), Inches(0.68), GREEN, line=GREEN)
    add_text(slide, "L'app couvre tout le cycle : chercher, choisir, gagner, suivre.", Inches(8.7), Inches(6.08), Inches(3.8), Inches(0.22), 12, WHITE, True, PP_ALIGN.CENTER)


def build_security_calc(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_title(slide, "Comment La Securite Est Calculee", "Le score final est sur 100.")
    add_box(slide, Inches(0.7), Inches(1.35), Inches(12.0), Inches(1.05), GREEN_SOFT, line=GREEN_SOFT)
    add_text(
        slide,
        "Score = (eclairage x 30%) + ((1 - trafic) x 25%) + (frequentation x 20%) + (piste cyclable x 25%)",
        Inches(0.95),
        Inches(1.65),
        Inches(11.4),
        Inches(0.35),
        18,
        GREEN_DARK,
        True,
        PP_ALIGN.CENTER,
    )
    items = [
        ("Eclairage", "30%", "Plus il y a de lampadaires, meilleur est le score.", ORANGE_SOFT, ORANGE),
        ("Trafic", "25%", "Un trafic fort reduit la securite.", BLUE_SOFT, BLUE),
        ("Frequentation", "20%", "Une zone plus vivante est souvent plus rassurante.", GREEN_SOFT, GREEN_DARK),
        ("Piste cyclable", "25%", "Une vraie infrastructure velo augmente la securite.", RGBColor(240, 236, 255), RGBColor(112, 89, 215)),
    ]
    x = Inches(0.75)
    for title, value, desc, bg, accent in items:
        add_box(slide, x, Inches(2.75), Inches(2.9), Inches(2.05), bg, line=bg)
        add_text(slide, title, x + Inches(0.18), Inches(3.0), Inches(2.4), Inches(0.22), 15, DARK, True)
        add_text(slide, value, x + Inches(0.18), Inches(3.36), Inches(1.2), Inches(0.35), 24, accent, True, font_name="Aptos Display")
        add_text(slide, desc, x + Inches(0.18), Inches(3.8), Inches(2.45), Inches(0.7), 10, TEXT)
        x += Inches(3.0)
    add_box(slide, Inches(0.75), Inches(5.3), Inches(5.75), Inches(1.35), WHITE, line=BORDER)
    add_text(slide, "Penalites meteo", Inches(1.0), Inches(5.55), Inches(2.0), Inches(0.25), 16, DARK, True)
    add_text(slide, "Orage : -20\nPluie : -10\nNuit : -15", Inches(1.0), Inches(5.95), Inches(1.6), Inches(0.6), 15, RED, True)
    add_text(slide, "Le score est ensuite limite entre 0 et 100.", Inches(2.8), Inches(5.98), Inches(2.8), Inches(0.35), 12, MUTED)
    add_box(slide, Inches(6.8), Inches(5.3), Inches(5.75), Inches(1.35), WHITE, line=BORDER)
    add_text(slide, "Lecture simple", Inches(7.05), Inches(5.55), Inches(2.0), Inches(0.25), 16, DARK, True)
    add_text(slide, "90+ : tres sur\n75-89 : bon trajet\n60-74 : correct\n<60 : a eviter", Inches(7.05), Inches(5.95), Inches(2.1), Inches(0.65), 15, GREEN_DARK, True)


def build_security_sources(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_title(slide, "D'ou Vient Le Score", "Le calcul s'appuie sur plusieurs sources.")
    panels = [
        ("OpenStreetMap / Overpass", "Lampadaires\nPistes cyclables\nZones pietonnes\nAnalyse autour d'un rayon de 300 m", GREEN_SOFT, GREEN_DARK, Inches(0.75)),
        ("Meteo Open-Meteo", "Conditions en temps reel\nPluie, orage, neige\nHeure du jour pour la penalite nuit", BLUE_SOFT, BLUE, Inches(4.4)),
        ("Base communautaire", "Routes et statistiques\nSignalements / donnees en base\nHistorique des trajets", ORANGE_SOFT, ORANGE, Inches(8.05)),
    ]
    for title, body, bg, accent, x in panels:
        add_box(slide, x, Inches(1.65), Inches(3.15), Inches(3.95), bg, line=bg)
        add_text(slide, title, x + Inches(0.2), Inches(1.95), Inches(2.6), Inches(0.35), 16, accent, True)
        add_text(slide, body, x + Inches(0.2), Inches(2.45), Inches(2.65), Inches(1.8), 13, TEXT)
    add_box(slide, Inches(1.2), Inches(6.05), Inches(10.95), Inches(0.8), WHITE, line=BORDER)
    add_text(slide, "Idee cle : l'application ne donne pas juste un itineraire, elle donne un contexte de confiance pour aider l'utilisateur a choisir.", Inches(1.45), Inches(6.28), Inches(10.5), Inches(0.28), 13, DARK, True, PP_ALIGN.CENTER)


def build_co2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_title(slide, "Comment Le CO2 Est Calcule", "Chaque mode est converti en emission totale pour le trajet.")
    add_box(slide, Inches(0.8), Inches(1.45), Inches(4.0), Inches(1.0), GREEN_SOFT, line=GREEN_SOFT)
    add_text(slide, "CO2 total = distance x facteur du mode", Inches(1.05), Inches(1.8), Inches(3.5), Inches(0.3), 20, GREEN_DARK, True, PP_ALIGN.CENTER)
    add_text(slide, "Exemple utilise dans l'ecran Transport", Inches(1.45), Inches(2.15), Inches(2.7), Inches(0.2), 11, MUTED, False, PP_ALIGN.CENTER)

    modes = [
        ("Velo", "0", GREEN, Inches(0.95), 0.0),
        ("Metro", "0.10", BLUE, Inches(2.0), 0.22),
        ("Bus", "0.16", ORANGE, Inches(3.05), 0.35),
        ("Voiture", "0.46", RED, Inches(4.1), 1.0),
    ]
    add_text(slide, "Facteur CO2 / km", Inches(0.95), Inches(2.8), Inches(3.0), Inches(0.25), 14, DARK, True)
    for label, value, color, x, ratio in modes:
        add_text(slide, label, x, Inches(3.2), Inches(0.8), Inches(0.2), 11, TEXT, True, PP_ALIGN.CENTER)
        bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + Inches(0.18), Inches(3.55), Inches(0.42), Inches(1.7 * ratio))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.color.rgb = color
        add_text(slide, value, x, Inches(5.45), Inches(0.8), Inches(0.2), 11, color, True, PP_ALIGN.CENTER)

    add_box(slide, Inches(5.25), Inches(1.45), Inches(7.0), Inches(2.2), WHITE, line=BORDER)
    add_text(slide, "Lecture du resultat", Inches(5.55), Inches(1.8), Inches(2.0), Inches(0.25), 18, DARK, True)
    add_text(slide, "Le mode le plus propre gagne plus de points eco.\nLa voiture sert de reference pour mesurer le CO2 evite.", Inches(5.55), Inches(2.2), Inches(5.9), Inches(0.7), 14, TEXT)
    add_text(slide, "Dans le comparateur, les options sont ensuite triees par points totaux.", Inches(5.55), Inches(3.0), Inches(5.6), Inches(0.3), 12, MUTED)

    add_picture_contain(slide, SHOTS / "02_transport.png", Inches(7.9), Inches(3.9), Inches(3.7), Inches(2.95))


def build_points(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_title(slide, "Comment Les Points Sont Calcules", "Le systeme combine l'effort ecologique et la securite.")
    add_box(slide, Inches(0.7), Inches(1.55), Inches(5.8), Inches(4.9), WHITE, line=BORDER)
    add_text(slide, "1. Points eco", Inches(1.0), Inches(1.9), Inches(2.0), Inches(0.25), 18, GREEN_DARK, True)
    add_text(slide, "ratio eco = (CO2 voiture - CO2 mode) / CO2 voiture", Inches(1.0), Inches(2.28), Inches(4.9), Inches(0.25), 14, TEXT)
    add_text(slide, "points eco = ratio eco x 25", Inches(1.0), Inches(2.65), Inches(3.3), Inches(0.25), 16, GREEN_DARK, True)
    add_text(slide, "Un trajet tres propre peut donc rapporter jusqu'a 25 points eco.", Inches(1.0), Inches(3.0), Inches(4.9), Inches(0.35), 12, MUTED)
    add_box(slide, Inches(0.95), Inches(3.7), Inches(5.2), Inches(1.95), GREEN_SOFT, line=GREEN_SOFT)
    add_text(slide, "Exemple", Inches(1.18), Inches(3.95), Inches(1.0), Inches(0.2), 15, DARK, True)
    add_text(slide, "Si le velo evite presque tout le CO2 par rapport a la voiture,\nil gagne presque 25 points eco.", Inches(1.18), Inches(4.35), Inches(4.7), Inches(0.6), 15, GREEN_DARK)

    add_box(slide, Inches(6.8), Inches(1.55), Inches(5.8), Inches(4.9), WHITE, line=BORDER)
    add_text(slide, "2. Points securite", Inches(7.1), Inches(1.9), Inches(2.5), Inches(0.25), 18, BLUE, True)
    add_text(slide, "90+  -> 15 pts\n75-89 -> 8 pts\n60-74 -> 5 pts\n<60   -> 0 pt", Inches(7.1), Inches(2.35), Inches(2.5), Inches(1.2), 18, BLUE, True)
    add_text(slide, "3. Total affiche dans l'ecran Transport", Inches(7.1), Inches(4.0), Inches(3.6), Inches(0.25), 18, ORANGE, True)
    add_text(slide, "points totaux = points eco + points securite", Inches(7.1), Inches(4.38), Inches(4.0), Inches(0.25), 15, TEXT)
    add_text(slide, "Exemple velo : 25 eco + 15 securite = 40 points", Inches(7.1), Inches(4.78), Inches(4.5), Inches(0.25), 16, ORANGE, True)
    add_text(slide, "Note : dans l'historique Firebase, un score resume 50/50 est aussi sauvegarde pour les stats.", Inches(7.1), Inches(5.45), Inches(4.8), Inches(0.4), 11, MUTED)


def build_architecture_team(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_title(slide, "Architecture Et Equipe", "Une repartition claire pour avancer vite a 4.")
    add_box(slide, Inches(0.7), Inches(1.45), Inches(12.0), Inches(1.1), BLUE_SOFT, line=BLUE_SOFT)
    add_text(slide, "Flutter (UI)  ->  Services metier  ->  APIs externes  ->  Firebase", Inches(1.15), Inches(1.83), Inches(11.0), Inches(0.25), 20, BLUE, True, PP_ALIGN.CENTER)
    roles = [
        ("Membre 1", "Carte\nNavigation\nMap\nItineraires", GREEN_SOFT, GREEN_DARK, Inches(0.85)),
        ("Membre 2", "Transport\nCO2\nPoints\nClassement", BLUE_SOFT, BLUE, Inches(3.95)),
        ("Membre 3", "Securite\nProfil\nStats\nAnalyse", ORANGE_SOFT, ORANGE, Inches(7.05)),
        ("Membre 4", "Firebase\nBoutique\nIntegration\nFinalisation", RGBColor(238, 247, 255), RGBColor(73, 107, 177), Inches(10.15)),
    ]
    for title, body, bg, accent, x in roles:
        add_box(slide, x, Inches(3.0), Inches(2.3), Inches(2.7), bg, line=bg)
        add_text(slide, title, x + Inches(0.16), Inches(3.25), Inches(1.8), Inches(0.25), 16, accent, True)
        add_text(slide, body, x + Inches(0.16), Inches(3.72), Inches(1.85), Inches(1.35), 15, TEXT)
    add_box(slide, Inches(0.85), Inches(6.15), Inches(11.95), Inches(0.65), GREEN, line=GREEN)
    add_text(slide, "Astuce soutenance : remplacez simplement Membre 1, 2, 3 et 4 par vos vrais noms avant de presenter.", Inches(1.1), Inches(6.34), Inches(11.4), Inches(0.22), 12, WHITE, True, PP_ALIGN.CENTER)


def build_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, DARK)
    add_text(slide, "EcoSafe", Inches(0.8), Inches(0.9), Inches(3.0), Inches(0.5), 28, WHITE, True, font_name="Aptos Display")
    add_text(slide, "Ce que l'on montre en demo", Inches(0.8), Inches(1.65), Inches(3.6), Inches(0.35), 18, WHITE, True)
    add_text(slide, "1. Recherche d'un trajet\n2. Comparaison des modes\n3. Lecture du score securite\n4. Gain de points\n5. Echange dans la boutique", Inches(0.8), Inches(2.2), Inches(3.8), Inches(1.8), 16, RGBColor(220, 229, 241))
    add_text(slide, "Merci", Inches(0.8), Inches(5.4), Inches(2.5), Inches(0.45), 26, GREEN, True, font_name="Aptos Display")
    add_text(slide, "Questions ?", Inches(0.8), Inches(5.95), Inches(2.5), Inches(0.35), 18, WHITE, True)
    add_picture_contain(slide, SHOTS / "05_boutique.png", Inches(5.8), Inches(0.8), Inches(2.8), Inches(5.9))
    add_picture_contain(slide, SHOTS / "02_transport.png", Inches(8.8), Inches(1.2), Inches(2.8), Inches(5.5))


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    build_cover(prs)
    build_problem_solution(prs)
    build_flow(prs)
    build_pages(prs)
    build_security_calc(prs)
    build_security_sources(prs)
    build_co2(prs)
    build_points(prs)
    build_architecture_team(prs)
    build_closing(prs)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
